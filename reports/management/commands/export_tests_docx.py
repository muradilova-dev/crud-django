"""
Команда: python manage.py export_tests_docx

Находит все тесты проекта, запускает их средствами Django
и записывает в Word-файл: таблицу результатов и исходный код тестов.
"""
import inspect
import io
import platform
import time
import unittest
from datetime import datetime

import django
from django.core.management.base import BaseCommand
from django.test.runner import DiscoverRunner
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

FONT = 'Times New Roman'
CODE_FONT = 'Courier New'


class RecordingResult(unittest.TextTestResult):
    """Запоминает результат и время каждого теста."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.records = []
        self._t0 = 0

    def startTest(self, test):
        super().startTest(test)
        self._t0 = time.perf_counter()

    def _add(self, test, status, details=''):
        self.records.append({
            'test': test,
            'status': status,
            'time': time.perf_counter() - self._t0,
            'details': details,
        })

    def addSuccess(self, test):
        super().addSuccess(test)
        self._add(test, 'OK')

    def addFailure(self, test, err):
        super().addFailure(test, err)
        self._add(test, 'FAIL', self._exc_info_to_string(err, test))

    def addError(self, test, err):
        super().addError(test, err)
        self._add(test, 'ERROR', self._exc_info_to_string(err, test))

    def addSkip(self, test, reason):
        super().addSkip(test, reason)
        self._add(test, 'SKIP', reason)


def test_description(test):
    """Описание теста: комментарий над методом или docstring."""
    func = getattr(test, test._testMethodName)
    comment = inspect.getcomments(func)
    if comment:
        return ' '.join(line.lstrip('# ').strip() for line in comment.strip().splitlines())
    doc = inspect.getdoc(func)
    return doc.splitlines()[0] if doc else ''


def set_cell_bg(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tc_pr.append(shd)


def set_run_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, size=12, bold=False, align=None, space_after=6):
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), FONT, size, bold)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), FONT, 16 if level == 1 else 14, True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    return p


def add_code(doc, source):
    for line in source.rstrip().splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line if line else ' ')
        set_run_font(run, CODE_FONT, 9)
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'F2F2F2')
        ppr.append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


class Command(BaseCommand):
    help = 'Запускает все тесты проекта и записывает их (код и результаты) в Word-файл'

    def add_arguments(self, parser):
        parser.add_argument('labels', nargs='*', help='например: blog crud (по умолчанию все приложения)')
        parser.add_argument('-o', '--output', default='tests_report.docx', help='имя выходного файла')

    # ---------- запуск тестов ----------
    def run_tests(self, labels):
        runner = DiscoverRunner(verbosity=0, interactive=False)
        suite = runner.build_suite(labels or None)
        runner.setup_test_environment()
        old_config = runner.setup_databases()
        started = time.perf_counter()
        try:
            result = unittest.TextTestRunner(
                stream=io.StringIO(), resultclass=RecordingResult, verbosity=0
            ).run(suite)
        finally:
            runner.teardown_databases(old_config)
            runner.teardown_test_environment()
        return result, time.perf_counter() - started

    @staticmethod
    def source_order(rec):
        """Сортировка: приложение, затем порядок классов и методов в файле."""
        test = rec['test']
        cls = test.__class__
        func = getattr(test, test._testMethodName)
        return (cls.__module__, inspect.getsourcelines(cls)[1], inspect.getsourcelines(func)[1])

    # ---------- сборка документа ----------
    def build_document(self, result, total_time):
        records = sorted(result.records, key=self.source_order)
        passed = sum(r['status'] == 'OK' for r in records)
        failed = sum(r['status'] in ('FAIL', 'ERROR') for r in records)

        doc = Document()
        section = doc.sections[0]
        section.left_margin, section.right_margin = Cm(2.5), Cm(1.5)
        section.top_margin = section.bottom_margin = Cm(2)

        add_text(doc, 'Отчёт о тестировании Django-проекта', 20, True, WD_ALIGN_PARAGRAPH.CENTER, 12)
        add_text(doc, f'Дата: {datetime.now():%d.%m.%Y %H:%M}')
        add_text(doc, f'Django {django.get_version()}, Python {platform.python_version()}')
        add_text(doc, f'Всего тестов: {len(records)}. Успешно: {passed}. Провалено: {failed}. '
                      f'Время выполнения: {total_time:.3f} с.')
        status_text = 'ВСЕ ТЕСТЫ ВЫПОЛНЕНЫ УСПЕШНО (OK)' if result.wasSuccessful() else 'ЕСТЬ ПРОВАЛЕННЫЕ ТЕСТЫ (FAILED)'
        add_text(doc, f'Итог: {status_text}', 13, True)

        # --- таблица результатов ---
        add_heading(doc, '1. Результаты тестов')
        widths = [Cm(0.9), Cm(3.3), Cm(4.6), Cm(4.8), Cm(1.5), Cm(1.6)]
        table = doc.add_table(rows=1, cols=6)
        table.style = 'Table Grid'
        table.autofit = False
        for i, w in enumerate(widths):
            table.columns[i].width = w
        for i, title in enumerate(['№', 'Класс', 'Тест', 'Что проверяет', 'Итог', 'Время, с']):
            cell = table.rows[0].cells[i]
            cell.width = widths[i]
            set_cell_bg(cell, 'D9E2F3')
            set_run_font(cell.paragraphs[0].add_run(title), FONT, 10, True)
        colors = {'OK': 'E2F0D9', 'FAIL': 'F8CBAD', 'ERROR': 'F8CBAD', 'SKIP': 'FFF2CC'}
        for n, rec in enumerate(records, 1):
            test = rec['test']
            row = table.add_row().cells
            values = [str(n), f"{test.__class__.__module__.split('.')[0]}.{test.__class__.__name__}",
                      test._testMethodName, test_description(test), rec['status'], f"{rec['time']:.3f}"]
            for i, value in enumerate(values):
                row[i].width = widths[i]
                set_run_font(row[i].paragraphs[0].add_run(value), FONT, 9, bold=(i == 4))
            set_cell_bg(row[4], colors.get(rec['status'], 'FFFFFF'))

        # --- код тестов (по классам) ---
        add_heading(doc, '2. Исходный код тестов')
        seen = []
        for rec in records:
            cls = rec['test'].__class__
            if cls not in seen:
                seen.append(cls)
        for cls in seen:
            add_heading(doc, f"{cls.__module__}.{cls.__name__}", 2)
            if cls.__doc__:
                add_text(doc, inspect.getdoc(cls).splitlines()[0], 11)
            add_code(doc, inspect.getsource(cls))

        # --- ошибки ---
        bad = [r for r in records if r['status'] in ('FAIL', 'ERROR')]
        if bad:
            add_heading(doc, '3. Ошибки')
            for rec in bad:
                add_text(doc, f"{rec['test'].id()} — {rec['status']}", 11, True)
                add_code(doc, rec['details'])
        return doc

    def handle(self, *args, **options):
        result, total_time = self.run_tests(options['labels'])
        doc = self.build_document(result, total_time)
        doc.save(options['output'])
        ok = result.testsRun - len(result.failures) - len(result.errors)
        self.stdout.write(self.style.SUCCESS(
            f"Готово: {options['output']} (тестов: {result.testsRun}, успешно: {ok})"))
